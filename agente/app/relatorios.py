# INTRODUÇÃO
# Relatórios gerados pelo agente: PDF (reportlab + matplotlib) e Excel (openpyxl).
# Decisões:
#   * reportlab em vez de WeasyPrint: não depende de pango/cairo do sistema e
#     cabe melhor no limite de 384 MB do container (justificativa no README);
#   * gráficos do PDF são PNGs gerados com matplotlib (Agg, sem display);
#   * arquivos têm nome aleatório (uuid) e são apagados após 24 h — a limpeza
#     roda a cada geração, o que dispensa cron dentro do container;
#   * os links devolvidos são relativos (/arquivos/<nome>); o frontend os
#     transforma em URL de download via /api/agente/arquivo/<nome>.
# RESUMO: gerar_pdf(titulo, periodo, analise, secoes) e gerar_excel(nome,
# colunas, linhas) gravam em ARQUIVOS_DIR e devolvem o nome do arquivo.

from __future__ import annotations

import os
import re
import time
import uuid
from decimal import Decimal
from html import escape
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # backend sem display (container não tem interface gráfica)
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ARQUIVOS_DIR = Path(os.environ.get("ARQUIVOS_DIR", "/app/arquivos_gerados"))
VALIDADE_SEGUNDOS = 24 * 3600  # arquivos gerados vivem no máximo 24 h


def _novo_nome(extensao: str) -> str:
    """Nome aleatório (nunca vem do usuário/modelo: evita path traversal)."""
    return f"{uuid.uuid4().hex}.{extensao}"


def limpar_antigos() -> None:
    """Apaga arquivos gerados há mais de 24 h (chamado a cada nova geração)."""
    ARQUIVOS_DIR.mkdir(parents=True, exist_ok=True)
    agora = time.time()
    for arquivo in ARQUIVOS_DIR.iterdir():
        if arquivo.is_file() and agora - arquivo.stat().st_mtime > VALIDADE_SEGUNDOS:
            arquivo.unlink(missing_ok=True)


def _grafico_png(grafico: dict, caminho: Path) -> None:
    """Desenha um gráfico simples (barra, linha ou pizza) em PNG para o PDF."""
    tipo = grafico.get("tipo", "barra")
    rotulos = [str(r) for r in grafico.get("rotulos", [])]
    valores = [float(v) for v in grafico.get("valores", [])]
    fig, ax = plt.subplots(figsize=(7, 3.2), dpi=110)
    if tipo == "pizza":
        ax.pie(valores, labels=rotulos, autopct="%1.1f%%")
    elif tipo == "linha":
        ax.plot(rotulos, valores, marker="o")
        ax.tick_params(axis="x", rotation=45)
    else:  # barra
        ax.bar(rotulos, valores)
        ax.tick_params(axis="x", rotation=45)
    if tipo != "pizza":
        # Eixo em padrão brasileiro (12.000.000) em vez de notação científica.
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: formatar(int(v))))
    if grafico.get("titulo"):
        ax.set_title(grafico["titulo"])
    fig.tight_layout()
    fig.savefig(caminho)
    plt.close(fig)


def formatar(valor) -> str:
    """Número em padrão brasileiro (1.234.567,89); demais valores como texto."""
    if isinstance(valor, bool) or valor is None:
        return "" if valor is None else str(valor)
    if isinstance(valor, int):
        return f"{valor:,}".replace(",", ".")
    if isinstance(valor, (float, Decimal)):
        if float(valor).is_integer() and abs(float(valor)) < 1e6:
            return formatar(int(valor))  # contagens somadas chegam como Decimal
        texto = f"{float(valor):,.2f}"
        return texto.replace(",", "X").replace(".", ",").replace("X", ".")
    return str(valor)


def _paragrafos(texto: str) -> list[str]:
    """Quebra a análise em parágrafos e converte **negrito** do markdown."""
    blocos = [b.strip() for b in re.split(r"\n\s*\n", texto or "") if b.strip()]
    saida = []
    for bloco in blocos:
        bloco = escape(bloco.lstrip("#").strip()).replace("\n", "<br/>")
        saida.append(re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", bloco))
    return saida


def gerar_pdf(titulo: str, periodo: str, analise: str,
              secoes: list[dict]) -> str:
    """Gera um relatório PDF e devolve o nome do arquivo.

    `analise` é o texto escrito pelo agente (resumo e conclusões).
    `secoes` é uma lista de blocos:
      {"tipo": "texto",   "texto": "..."}
      {"tipo": "tabela",  "titulo": "...", "nota": "...", "colunas": [...],
                          "linhas": [[...]]}
      {"tipo": "grafico", "grafico": {"tipo": "barra|linha|pizza",
                                      "titulo": "...", "rotulos": [...],
                                      "valores": [...]}}
    """
    limpar_antigos()
    nome = _novo_nome("pdf")
    caminho = ARQUIVOS_DIR / nome

    estilos = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(caminho), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    story = [
        Paragraph(escape(titulo), estilos["Title"]),
        Paragraph(escape(f"Período: {periodo}"), estilos["Normal"]),
        Paragraph("Dados fictícios gerados para fins de demonstração "
                  "(Serra Clara Imóveis).", estilos["Italic"]),
        Spacer(1, 0.5 * cm),
        Paragraph("Análise", estilos["Heading2"]),
    ]
    for paragrafo in _paragrafos(analise):
        story += [Paragraph(paragrafo, estilos["Normal"]), Spacer(1, 0.25 * cm)]
    story.append(Spacer(1, 0.3 * cm))

    for secao in secoes:
        tipo = secao.get("tipo")
        if tipo == "texto":
            story.append(Paragraph(escape(str(secao.get("texto", ""))),
                                   estilos["Normal"]))
            story.append(Spacer(1, 0.3 * cm))
        elif tipo == "tabela":
            if secao.get("titulo"):
                story.append(Paragraph(escape(str(secao["titulo"])), estilos["Heading3"]))
            if secao.get("nota"):
                story.append(Paragraph(escape(str(secao["nota"])), estilos["Italic"]))
            colunas = [str(c) for c in secao.get("colunas", [])]
            linhas = [[formatar(v) for v in linha]
                      for linha in secao.get("linhas", [])]
            cabecalho = [Paragraph(f'<font color="white"><b>{escape(c)}</b></font>',
                                   estilos["Normal"]) for c in colunas]
            tabela = Table([cabecalho, *[[Paragraph(escape(v), estilos["Normal"]) for v in row]
                                         for row in linhas]],
                           colWidths=[17 * cm / max(len(colunas), 1)] * len(colunas), repeatRows=1)
            tabela.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9ca3af")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#f3f4f6")]),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            story.append(tabela)
            story.append(Spacer(1, 0.4 * cm))
        elif tipo == "grafico":
            png = ARQUIVOS_DIR / f"{uuid.uuid4().hex}.png"
            try:
                _grafico_png(secao.get("grafico", {}), png)
                story.append(Image(BytesIO(png.read_bytes()), width=16 * cm, height=7.3 * cm))
                story.append(Spacer(1, 0.4 * cm))
            finally:
                png.unlink(missing_ok=True)  # o PNG é só insumo do PDF

    doc.build(story)
    return nome


def gerar_excel(nome_planilha: str, colunas: list[str],
                linhas: list[list]) -> str:
    """Gera uma planilha .xlsx com uma aba e devolve o nome do arquivo."""
    limpar_antigos()
    nome = _novo_nome("xlsx")
    caminho = ARQUIVOS_DIR / nome

    wb = Workbook()
    ws = wb.active
    ws.title = (nome_planilha or "dados")[:31]  # limite do Excel para abas
    ws.append([str(c) for c in colunas])
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        ws.append(list(linha))
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # valores externos nunca viram fórmulas
    for i, _ in enumerate(colunas, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
    wb.save(caminho)
    return nome
